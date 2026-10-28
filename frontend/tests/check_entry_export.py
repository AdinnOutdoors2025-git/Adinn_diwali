"""Independently check values, theme, source BD names, and native Excel chart references."""
import json
import math
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils.cell import range_boundaries

filename, expected_file = sys.argv[1:3]
with open(expected_file, encoding='utf-8') as stream:
    expected = json.load(stream)
workbook = openpyxl.load_workbook(filename)
assert workbook.sheetnames == ['Summary', 'Distribution Details', 'Type Analytics', 'Requestee Analytics']
summary, details, type_sheet, bd_sheet = [workbook[name] for name in workbook.sheetnames]
navy, orange, ivory, white, green, rose = 'FF14213D', 'FFF59E0B', 'FFFFFBEB', 'FFFFFFFF', 'FF059669', 'FFE11D48'
assert summary['A1'].value == 'ADINN | Diwali Sweet Distribution Report'
assert summary['A1'].fill.fgColor.rgb == navy
assert summary['A1'].font.color.rgb == white and summary['A1'].font.bold
assert summary.freeze_panes == 'A8'
assert 'A1:L2' in [str(value) for value in summary.merged_cells]
assert summary['B3'].value == expected['filters']['year']
assert summary['B4'].value == ('All Types' if expected['filters']['type'] == 'ALL' else expected['filters']['type'])
assert summary['B5'].value == (expected['filters']['search'] or 'All clients')
assert summary['B6'].value == {'ALL': 'All', 'GIVEN': 'Given', 'NOT_GIVEN': 'Not Given'}[expected['filters']['status']]
assert summary['A8'].font.color.rgb == navy
records = expected['records']
bds = expected['bds']
given = sum(not row['isEligible'] for row in records)
pending = len(records) - given
percent = lambda given, registered: math.floor(given / registered * 10000 + 0.5) / 10000 if registered else 0
assert [summary[address].value for address in ['A9', 'C9', 'E9', 'G9']] == [len(records), given, given, pending]
assert math.isclose(summary['I9'].value, percent(given, len(records)), abs_tol=1e-10)
assert summary['I9'].number_format == '0.00%'

assert [cell.value for cell in details[1]] == expected['headings']
assert details['K1'].value == 'Distribution Marked'
assert details.max_row == len(expected['rows']) + 1
for index, wanted in enumerate(expected['rows'], start=2):
    actual = [cell.value for cell in details[index]]
    assert [value if value is not None else '' for value in actual] == wanted
    assert isinstance(details.cell(index, 3).value, str), 'Phones must retain text data type'
    # The pre-existing IST text values are preserved, not reinterpreted as a different instant.
    assert (details.cell(index, 11).value or '') == wanted[10]
    assert details.cell(index, 11).number_format == 'dd mmm yyyy hh:mm:ss'
    assert details.cell(index, 10).font.color.rgb == (green if wanted[9] == 'Given' else rose)
    assert details.cell(index, 2).fill.fgColor.rgb == (white if index % 2 == 0 else ivory)
    assert details.row_dimensions[index].height >= 32

for sheet in [details, type_sheet, bd_sheet]:
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref is not None
    for cell in sheet[1]:
        assert cell.fill.fgColor.rgb == navy and cell.font.color.rgb == white and cell.font.bold
        assert cell.border.bottom.color.rgb == orange
        assert cell.alignment.wrap_text
    assert sheet.row_dimensions[1].height >= 32
    assert sheet.page_setup.orientation == 'landscape'
    assert sheet.column_dimensions['A'].width >= 8
    for row in range(2, sheet.max_row + 1):
        assert sheet.cell(row, 1).fill.fgColor.rgb == (white if row % 2 == 0 else ivory)
for sheet in workbook:
    assert not any(cell.value == 'Distribution Date (IST)' for row in sheet for cell in row)

for row in range(2, type_sheet.max_row + 1):
    values = [cell.value for cell in type_sheet[row]]
    name = values[0]
    members = [record for record in records if any(link['type'] == name for link in record['associations'])]
    given_here = sum(not record['isEligible'] for record in members)
    credited = sum(record['currentYearDistribution'] is not None and record['currentYearDistribution']['type'] == name for record in records)
    assert values[:4] == [name, len(members), given_here, len(members) - given_here]
    assert math.isclose(values[4], percent(given_here, len(members)), abs_tol=1e-10)
    assert values[5] == credited
    assert type_sheet.cell(row, 5).number_format == '0.00%'

assert bd_sheet.max_row == len(bds) + 1
assert [(bd_sheet.cell(row, 1).value, bd_sheet.cell(row, 2).value) for row in range(2, bd_sheet.max_row + 1)] == [(bd['name'], bd['email']) for bd in bds]
for index, bd in enumerate(bds, start=2):
    members = [record for record in records if any(link['bd'] and link['bd']['_id'] == bd['_id'] for link in record['associations'])]
    given_here = sum(not record['isEligible'] for record in members)
    credited = sum(record['currentYearDistribution'] is not None and record['currentYearDistribution']['bdId'] == bd['_id'] for record in records)
    values = [cell.value for cell in bd_sheet[index]]
    assert values[2:6] == [len(members), given_here, len(members) - given_here, credited]
    assert math.isclose(values[6], percent(given_here, len(members)), abs_tol=1e-10)
    assert bd_sheet.cell(index, 7).number_format == '0.00%'

assert len(summary._charts) == 3
assert [type(chart).__name__ for chart in summary._charts] == ['DoughnutChart', 'BarChart', 'BarChart']
ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(filename) as archive:
    assert archive.testzip() is None
    charts = [name for name in archive.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml', name)]
    assert len(charts) == 3
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
    for name in charts:
        root = ET.fromstring(archive.read(name))
        for reference_type, cache_type, numeric in [('strRef', 'strCache', False), ('numRef', 'numCache', True)]:
            for reference in root.findall(f'.//c:{reference_type}', ns):
                sheet_name, cell_range = reference.find('c:f', ns).text.split('!')
                first_col, first_row, last_col, last_row = range_boundaries(cell_range.replace('$', ''))
                sheet = workbook[sheet_name.strip("'")]
                actual = [sheet.cell(row, col).value for row in range(first_row, last_row + 1) for col in range(first_col, last_col + 1)]
                actual = [value if value is not None else '' for value in actual]
                cached = [point.find('c:v', ns).text or '' for point in reference.findall(f'c:{cache_type}/c:pt', ns)]
                assert len(actual) == len(cached)
                assert ([float(value) for value in cached] == actual) if numeric else (cached == actual)
                if name.endswith('chart3.xml') and not numeric:
                    assert all(value in [bd['name'] for bd in bds] or not value for value in actual)
    colors = lambda number: [element.attrib['val'] for element in ET.fromstring(archive.read(f'xl/charts/chart{number}.xml')).findall('.//a:srgbClr', ns)]
    assert '059669' in colors(1) and 'E11D48' in colors(1)
    assert '14213D' in colors(2) and 'F59E0B' in colors(2)
    assert '14213D' in colors(3)
    if len(summary._charts[2].series[0].cat.strRef.strCache.pt) >= 2:
        assert 'FCD34D' in colors(3)
print(json.dumps({'file': filename, 'sheets': 4, 'rows': len(records), 'bd_entries': len(bds), 'charts': 3, 'theme_verified': True, 'values_and_timestamps_preserved': True, 'chart_ranges_and_colors_verified': True}))
